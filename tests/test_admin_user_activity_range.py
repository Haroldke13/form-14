import io
from datetime import date, datetime

from openpyxl import load_workbook

from test_security_and_workflows import create_user, login


def _seed_activity(app, models):
    db = models["db"]
    User = models["User"]
    PBOReport = models["PBOReport"]

    with app.app_context():
        from models import UserActivityLog

        create_user(db, User, "range-admin@example.com", "RangeAdmin@123", role="admin")
        early_user = create_user(db, User, "early-user@example.com", "EarlyUser@123")
        mid_user = create_user(db, User, "mid-user@example.com", "MidUser@123")
        late_user = create_user(db, User, "late-user@example.com", "LateUser@123")

        # Pin creation timestamps: a report created inside the window counts as a
        # submission event, so leaving the default (now) would leak rows into
        # open-ended ranges that run up to today.
        early_report = PBOReport(pbo_name="Early Report", user_id=early_user.id, created_at=datetime(2026, 3, 1, 8, 0, 0))
        mid_report = PBOReport(pbo_name="Mid Report", user_id=mid_user.id, created_at=datetime(2026, 3, 10, 8, 0, 0))
        late_report = PBOReport(pbo_name="Late Report", user_id=late_user.id, created_at=datetime(2026, 3, 25, 8, 0, 0))
        db.session.add_all([early_report, mid_report, late_report])
        db.session.commit()

        db.session.add_all([
            UserActivityLog(user_id=early_user.id, report_id=early_report.id, action="report_updated", created_at=datetime(2026, 3, 1, 9, 0, 0)),
            UserActivityLog(user_id=mid_user.id, report_id=mid_report.id, action="report_updated", created_at=datetime(2026, 3, 10, 12, 0, 0)),
            UserActivityLog(user_id=late_user.id, report_id=late_report.id, action="report_updated", created_at=datetime(2026, 3, 25, 16, 0, 0)),
        ])
        db.session.commit()


def test_resolve_admin_activity_date_range_handles_all_inputs(app, app_code):
    resolve = app_code.resolve_admin_activity_date_range

    with app.test_request_context("/admin"):
        result = resolve()
        assert result["is_active"] is False
        assert result["start"] is None and result["end"] is None
        assert result["label"] == "" and result["slug"] == "all"

    with app.test_request_context("/admin?activity_from=2026-03-05&activity_to=2026-03-15"):
        result = resolve()
        assert result["start"] == datetime(2026, 3, 5, 0, 0, 0)
        assert result["end"] == datetime.combine(date(2026, 3, 15), datetime.max.time())
        assert result["label"] == "2026-03-05 to 2026-03-15"
        assert result["slug"] == "20260305_20260315"

    with app.test_request_context("/admin?activity_from=2026-03-15&activity_to=2026-03-05"):
        result = resolve()
        assert result["from_value"] == "2026-03-05"
        assert result["to_value"] == "2026-03-15"

    with app.test_request_context("/admin?activity_from=2026-03-20"):
        result = resolve()
        assert result["start"] == datetime(2026, 3, 20, 0, 0, 0)
        assert result["end"] is not None and result["end"].date() >= date(2026, 3, 20)
        assert result["label"] == "from 2026-03-20"

    with app.test_request_context("/admin?activity_to=2026-03-05"):
        result = resolve()
        assert result["start"] == datetime.combine(app_code.ADMIN_ACTIVITY_RANGE_FLOOR, datetime.min.time())
        assert result["label"] == "up to 2026-03-05"

    with app.test_request_context("/admin?activity_date=2026-03-10"):
        result = resolve()
        assert result["from_value"] == "2026-03-10" and result["to_value"] == "2026-03-10"
        assert result["label"] == "2026-03-10"
        assert result["slug"] == "20260310"

    with app.test_request_context("/admin?activity_from=not-a-date"):
        result = resolve()
        assert result["is_active"] is False


def test_admin_user_activity_panel_filters_by_date_range(app, client, models):
    _seed_activity(app, models)
    login(client, "range-admin@example.com", "RangeAdmin@123")

    response = client.get("/admin/user-activity-panel?activity_from=2026-03-05&activity_to=2026-03-15")
    assert response.status_code == 200
    assert b"mid-user@example.com" in response.data
    assert b"early-user@example.com" not in response.data
    assert b"late-user@example.com" not in response.data

    open_ended = client.get("/admin/user-activity-panel?activity_from=2026-03-09")
    assert open_ended.status_code == 200
    assert b"mid-user@example.com" in open_ended.data
    assert b"late-user@example.com" in open_ended.data
    assert b"early-user@example.com" not in open_ended.data

    legacy = client.get("/admin/user-activity-panel?activity_date=2026-03-01")
    assert legacy.status_code == 200
    assert b"early-user@example.com" in legacy.data
    assert b"mid-user@example.com" not in legacy.data

    empty = client.get("/admin/user-activity-panel?activity_from=2026-04-01&activity_to=2026-04-30")
    assert empty.status_code == 200
    assert b"No user activity found for 2026-04-01 to 2026-04-30" in empty.data


def test_admin_user_activity_export_uses_range_label_and_slug(app, client, models):
    _seed_activity(app, models)
    login(client, "range-admin@example.com", "RangeAdmin@123")

    response = client.get("/admin/user-activity/export?activity_from=2026-03-05&activity_to=2026-03-15")
    assert response.status_code == 200
    assert "user_activity_20260305_20260315_" in response.headers["Content-Disposition"]

    worksheet = load_workbook(io.BytesIO(response.data))["User Activity"]
    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[0][0] == "Activity Period"
    assert rows[1][0] == "2026-03-05 to 2026-03-15"
    assert rows[1][1] == "mid-user@example.com"
    assert len(rows) == 2


def test_admin_dashboard_layout_after_review_queue_move(app, client, models):
    _seed_activity(app, models)
    login(client, "range-admin@example.com", "RangeAdmin@123")

    response = client.get("/admin?activity_from=2026-03-05&activity_to=2026-03-15")
    assert response.status_code == 200
    html = response.data.decode("utf-8")

    assert "Workflow Settings" not in html
    assert 'name="edit_deadline"' not in html
    assert "Review Queue" in html
    assert 'id="review-queue-slider"' in html
    assert 'id="activity_from"' in html and 'id="activity_to"' in html
    assert 'name="activity_date"' not in html
    assert "Showing activity 2026-03-05 to 2026-03-15" in html
    assert html.index("Review Queue") < html.index("Export Reports")
    assert html.index("Export Reports") < html.index("User Activity")
    assert "Backup to Google" in html and "Field Help Rules" in html
